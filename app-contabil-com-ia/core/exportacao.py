"""XLSX real com valores literais: texto importado nunca vira fórmula."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def planilha_tabela(colunas, linhas):
    book=Workbook();sheet=book.active;sheet.title='Relatório'
    sheet.append(colunas)
    for linha in linhas:sheet.append(linha)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s'
    for cell in sheet[1]:
        cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='FF7500')
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(45,max(14,max(len(str(c.value or '')) for c in col)+2))
    stream=BytesIO();book.save(stream);book.close();return stream.getvalue()

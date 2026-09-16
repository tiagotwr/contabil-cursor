"""Leitura pura de planilhas para a etapa de prévia da importação."""
from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from core.plano_contas import organizar_contas_com_avisos

try:
    from openpyxl import load_workbook
except ImportError:  # mensagem útil para quem usar a prévia fora do ambiente do app
    load_workbook = None


CAMPOS = [
    "linha_id", "documento_id", "empresa_id", "data", "conta_id",
    "debito_centavos", "credito_centavos", "tipo", "historico", "origem_id",
    "centro_id", "atividade_caixa", "rubrica_caixa",
]
MAX_ROWS = 5_000
MAX_XLSX_BYTES = 12 * 1024 * 1024
MAX_XLSX_EXPANDED = 64 * 1024 * 1024

NORMAL_LANCAMENTOS = {
    "linha_id", "documento_id", "data", "conta_id", "tipo",
    "historico", "centro_id", "atividade_caixa", "rubrica_caixa",
}
LEGADO_MOVIMENTOS = {
    "CT2_DATA", "CT2_DC", "CT2_DOC", "CT2_LINHA", "CT2_DEBITO", "CT2_CREDIT", "CT2_VALOR",
}


def _text(value: object, field: str, *, numeric: bool = False) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        return ""
    if isinstance(value, bool):
        raise ValueError(f"{field} inválido.")
    if isinstance(value, (int, float, Decimal)) and not numeric:
        if isinstance(value, float) and not value.is_integer():
            raise ValueError(f"{field} numérico é ambíguo.")
        return str(int(value))
    if numeric:
        if isinstance(value, bool) or (isinstance(value, float) and not value.is_integer()):
            raise ValueError(f"{field} numérico é ambíguo.")
        return str(int(value)) if isinstance(value, (int, float, Decimal)) else str(value).strip()
    return str(value).strip()


def _date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str):
        raise ValueError("Data deve ser data Excel, YYYY-MM-DD, YYYYMMDD ou DD/MM/YYYY.")
    value = value.strip()
    for fmt in ("%Y-%m-%d", "%Y%m%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError("Data inválida ou ambígua; use YYYY-MM-DD, YYYYMMDD ou DD/MM/YYYY.")


def _centavos(value: object) -> int:
    if value is None or value == "":
        return 0
    try:
        text = str(value).strip()
        # CSV brasileiro usa ponto para milhar e vírgula para decimal.
        amount = Decimal(text.replace(".", "").replace(",", ".") if "," in text else text)
    except (InvalidOperation, AttributeError):
        raise ValueError("Valor monetário inválido.")
    if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -2:
        raise ValueError("Valor monetário deve ser não negativo e ter no máximo duas casas.")
    return int(amount * 100)


def _defaults(raw: bytes, filename: str) -> tuple[str, str]:
    fingerprint = hashlib.sha256(raw).hexdigest()[:20]
    return "EMPRESA_PLANILHA", f"planilha:{filename.lower()}:{fingerprint}"


def _validate_row(row: dict) -> dict:
    if set(row) != set(CAMPOS):
        raise ValueError("Colunas internas inválidas.")
    if not row["linha_id"] or not row["documento_id"] or not row["conta_id"]:
        raise ValueError("Linha, documento e conta são obrigatórios.")
    if row["debito_centavos"] and row["credito_centavos"]:
        raise ValueError("Uma linha não pode ter débito e crédito juntos.")
    if not row["debito_centavos"] and not row["credito_centavos"]:
        raise ValueError("Linha sem débito ou crédito.")
    return row


def _normal_row(source: dict, empresa: str, origem: str, index: int) -> dict:
    linha = _text(source.get("linha_id"), "linha_id")
    if not linha:
        raise ValueError("linha_id é obrigatório no modelo normalizado.")
    debit = source.get("debito_centavos") if "debito_centavos" in source else source.get("debito")
    credit = source.get("credito_centavos") if "credito_centavos" in source else source.get("credito")
    row = {
        "linha_id": linha, "documento_id": _text(source.get("documento_id"), "documento_id"),
        "empresa_id": _text(source.get("empresa_id"), "empresa_id") or empresa,
        "data": _date(source.get("data")), "conta_id": _text(source.get("conta_id"), "conta_id"),
        "debito_centavos": int(debit) if "debito_centavos" in source and str(debit or "0").isdigit() else _centavos(debit),
        "credito_centavos": int(credit) if "credito_centavos" in source and str(credit or "0").isdigit() else _centavos(credit),
        "tipo": _text(source.get("tipo"), "tipo"), "historico": _text(source.get("historico"), "historico"),
        "origem_id": _text(source.get("origem_id"), "origem_id") or origem,
        "centro_id": _text(source.get("centro_id"), "centro_id"),
        "atividade_caixa": _text(source.get("atividade_caixa"), "atividade_caixa"),
        "rubrica_caixa": _text(source.get("rubrica_caixa"), "rubrica_caixa"),
    }
    return _validate_row(row)


def _legacy_rows(source: dict, empresa: str, origem: str) -> list[dict]:
    side = _text(source.get("CT2_DC"), "CT2_DC", numeric=True)
    if side not in {"1", "2", "3"}:
        raise ValueError("CT2_DC deve ser 1 (débito), 2 (crédito) ou 3 (ambos).")
    recno = _text(source.get("R_E_C_N_O_"), "R_E_C_N_O_", numeric=True)
    document = _text(source.get("APP_DOCUMENTO_ID") or source.get("_documento_normalizado", source.get("CT2_DOC")), "CT2_DOC", numeric=True)
    line = _text(source.get("CT2_LINHA"), "CT2_LINHA", numeric=True)
    if not document and recno:
        document = f"RECNO-{recno}"
    if not document or not line:
        raise ValueError("CT2_DOC ausente sem R_E_C_N_O_, ou CT2_LINHA ausente.")
    amount = _centavos(source.get("CT2_VALOR"))
    if not amount:
        raise ValueError("CT2_VALOR deve ser maior que zero.")
    common = {
        "documento_id": document, "empresa_id": _text(source.get('APP_EMPRESA_ID'), 'APP_EMPRESA_ID') or empresa, "data": _date(source.get("CT2_DATA")),
        "tipo": _text(source.get('APP_TIPO'), 'APP_TIPO') or "competencia", "historico": _text(source.get("CT2_HIST"), "CT2_HIST"), "origem_id": _text(source.get('APP_ORIGEM_ID'), 'APP_ORIGEM_ID') or origem,
    }
    sides = []
    if side in {"1", "3"}:
        sides.append(("D", "CT2_DEBITO", "CT2_CCD", "CT2_ATIVDE", "CT2_ITEMD", "debito_centavos"))
    if side in {"2", "3"}:
        sides.append(("C", "CT2_CREDIT", "CT2_CCC", "CT2_ATIVCR", "CT2_ITEMC", "credito_centavos"))
    rows = []
    for label, account, center, activity, item, money in sides:
        suffix = f":{recno}" if recno else ""
        app_line = 'APP_LINHA_DEBITO' if label == 'D' else 'APP_LINHA_CREDITO'
        row = dict(common, linha_id=_text(source.get(app_line), app_line) or f"{empresa}:{document}:{line}:{label}{suffix}", conta_id=_text(source.get(account), account),
                   centro_id=_text(source.get(center), center), atividade_caixa=_text(source.get('APP_ATIVIDADE_CAIXA'), 'APP_ATIVIDADE_CAIXA'),
                   rubrica_caixa=_text(source.get('APP_RUBRICA_CAIXA') if 'APP_RUBRICA_CAIXA' in source else source.get(item), item), debito_centavos=0, credito_centavos=0)
        row[money] = amount
        rows.append(_validate_row(row))
    return rows


_CLASSES_BP_MODELO_ANTERIOR = {"circulante", "nao_circulante", "pl", "resultado"}


def _ct1_analitica(
    record: dict, conta_id: str, app_classe_bp: str, app_analitica: bool | None
) -> tuple[bool | None, bool]:
    """Traduz a convenção CT1 comprovada na planilha original.

    CT1_CLASSE 1 é sintética e 2 é analítica. CONTA_TIPO/CT1_ZTIPOC é uma
    classificação gerencial e não participa desta decisão.
    """
    classe = _text(record.get("CT1_CLASSE"), "CT1_CLASSE")
    if not classe:
        return None, False
    if classe == "1":
        return False, False
    if classe == "2":
        return True, False
    classe_anterior = classe.casefold()
    if (
        classe_anterior in _CLASSES_BP_MODELO_ANTERIOR
        and classe_anterior == app_classe_bp.casefold()
        and app_analitica is not None
    ):
        return app_analitica, True
    raise ValueError(
        f"Conta {conta_id}: CT1_CLASSE deve ser 1 (sintética) ou 2 (analítica); "
        "o modelo anterior só é aceito com APP_CLASSE_BP idêntico e APP_ANALITICA válido."
    )


def _bool_explicito(value: object, conta_id: str, campo: str) -> bool | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().casefold()
    if text in {"1", "true", "sim", "s", "analitica", "analítica"}:
        return True
    if text in {"0", "false", "nao", "não", "n", "sintetica", "sintética"}:
        return False
    raise ValueError(f"Conta {conta_id}: {campo} deve informar analítica ou sintética.")


def _accounts(records: list[dict], legacy: bool, avisos: list[str] | None = None) -> list[dict]:
    from core.dre_importacao import destino_planilha
    avisos = [] if avisos is None else avisos
    accounts = []
    for record in records:
        account = _text(record.get("CT1_CONTA") if legacy else record.get("conta_id"), "conta_id")
        if not account:
            raise ValueError("Conta sem código no catálogo.")
        if legacy:
            app = {key.lower()[4:]: value for key, value in record.items() if key.startswith("APP_")}
            app_analitica = _bool_explicito(app.get("analitica"), account, "APP_ANALITICA")
            classe, modelo_anterior = _ct1_analitica(
                record, account, _text(app.get("classe_bp"), "APP_CLASSE_BP"), app_analitica
            )
            if modelo_anterior:
                aviso = "Modelo Protheus anterior reconhecido: CT1_CLASSE patrimonial foi confirmado pelos campos APP_* explícitos."
                if aviso not in avisos:
                    avisos.append(aviso)
            if classe is not None and app_analitica is not None and classe != app_analitica:
                raise ValueError(f"Conta {account}: CT1_CLASSE conflita com APP_ANALITICA.")
            pai_ct1 = _text(record.get("CT1_CTASUP"), "CT1_CTASUP")
            pai_app = _text(app.get("conta_pai_id"), "APP_CONTA_PAI_ID")
            if pai_ct1 and pai_app and pai_ct1 != pai_app:
                raise ValueError(f"Conta {account}: CT1_CTASUP conflita com APP_CONTA_PAI_ID.")
            source = dict(app)
            source.update(conta_id=account, descricao=record.get("CT1_DESC01"))
            source["conta_pai_id"] = pai_app or pai_ct1
            if classe is not None:
                source["analitica"] = classe
            source["metadados_origem"] = {
                key: str(value) for key, value in record.items()
                if key not in {"CT1_CONTA", "CT1_DESC01"}
            }
        else:
            source = record
        accounts.append({
            "conta_id": account,
            "descricao": _text(source.get("descricao"), "descricao"),
            "grupo": _text(source.get("grupo"), "grupo") or "pendente",
            "componente": _text(source.get("componente"), "componente"),
            "linha_dre": _text(source.get("linha_dre"), "linha_dre"),
            "conta_pai_id": _text(source.get("conta_pai_id"), "conta_pai_id"),
            "ordem": _text(source.get("ordem"), "ordem", numeric=True),
            "nivel": _text(source.get("nivel"), "nivel", numeric=True),
            "natureza": _text(source.get("natureza"), "natureza"),
            "classe_bp": _text(source.get("classe_bp"), "classe_bp") or "pendente",
            "metadados_origem": source.get("metadados_origem", {}),
            # None mantém a ausência de tipo para o organizador decidir apenas
            # pela estrutura completa, sem supor que toda conta CT1 é folha.
            "analitica": source.get("analitica"),
        })
        if 'dre_destino' in source:
            accounts[-1]['dre_destino'] = destino_planilha(source['dre_destino'])
    return accounts


def _read_csv(raw: bytes) -> tuple[dict[str, list[dict]], set[str]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ValueError("CSV deve estar em UTF-8.")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;")
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
        raise ValueError("Cabeçalho CSV inválido ou repetido.")
    rows = list(reader)
    if not rows or len(rows) > MAX_ROWS or any(None in row or None in row.values() for row in rows):
        raise ValueError("CSV vazio, grande demais ou com colunas inconsistentes.")
    return {"Lançamentos": rows}, set(reader.fieldnames)


def _read_xlsx(raw: bytes) -> dict[str, list[dict]]:
    if load_workbook is None:
        raise ValueError("A leitura XLSX requer a dependência openpyxl.")
    if len(raw) > MAX_XLSX_BYTES:
        raise ValueError("XLSX excede o limite de tamanho.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            sizes = sum(item.file_size for item in archive.infolist())
            if len(archive.infolist()) > 2_000 or sizes > MAX_XLSX_EXPANDED or sizes > max(len(raw), 1) * 100:
                raise ValueError("XLSX excede os limites de expansão.")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        cached = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("XLSX inválido.") from error
    sheets, avisos = {}, []
    try:
        for name in ("Lançamentos", "Contas", "MOVIMENTOS", "CONTAS", "CENCUS", "DRE_GERENCIAL", "ORCAMENTO"):
            if name not in book.sheetnames:
                continue
            sheet, cached_sheet = book[name], cached[name]
            header_row = None
            for candidate in range(1, 21):
                cells = next(sheet.iter_rows(min_row=candidate, max_row=candidate), ())
                headers = [str(cell.value).strip() if cell.value is not None else "" for cell in cells]
                required = LEGADO_MOVIMENTOS if name == "MOVIMENTOS" else NORMAL_LANCAMENTOS if name == "Lançamentos" else {"CT1_CONTA"} if name == "CONTAS" else {"CTT_CUSTO"} if name == "CENCUS" else {'codigo','nome','ordem','tipo'} if name == 'DRE_GERENCIAL' else {"competencia","conta_id","valor_dc_centavos"} if name == "ORCAMENTO" else {"conta_id"}
                if required <= set(headers):
                    header_row = candidate
                    break
            if header_row is None:
                if name in ('DRE_GERENCIAL', 'ORCAMENTO'):
                    campos = 'codigo, nome, ordem, tipo' if name == 'DRE_GERENCIAL' else 'competencia, conta_id, valor_dc_centavos'
                    raise ValueError(f'{name}: cabeçalho obrigatório {campos}.')
                continue
            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row))]
            if len([h for h in headers if h]) != len(set(h for h in headers if h)):
                raise ValueError(f'Cabeçalho repetido em {name}.')
            records = []
            for cells, cached_cells in zip(sheet.iter_rows(min_row=header_row + 1), cached_sheet.iter_rows(min_row=header_row + 1)):
                record = {header: cell.value for header, cell in zip(headers, cells) if header}
                # Um rodapé não é fato se data, documento e ambas contas estiverem vazios.
                if name == "MOVIMENTOS" and not any(record.get(key) for key in ("CT2_DATA", "CT2_DOC", "CT2_DEBITO", "CT2_CREDIT")):
                    if any(cell.data_type == "f" for cell in cells):
                        avisos.append("linhasignoradas: rodapé sem campos de fato.")
                    continue
                if not any(value is not None for value in record.values()):
                    continue
                for header, cell, cached_cell in zip(headers, cells, cached_cells):
                    if header and cell.data_type == "f":
                        if cached_cell.value is None or not isinstance(cached_cell.value, (int, float)):
                            raise ValueError(f"Fórmula sem valor pré-calculado em {name}/{header}.")
                        record[header] = cached_cell.value
                        avisos.append("valor pré-calculado Excel; não recalculado app")
                record['__linha_excel__'] = cells[0].row
                records.append(record)
                if len(records) > MAX_ROWS:
                    raise ValueError("Planilha excede o limite de linhas.")
            sheets[name] = records
    finally:
        book.close(); cached.close()
    sheets["__avisos__"] = list(dict.fromkeys(avisos))
    return sheets


def ler_planilha(raw: bytes, filename: str) -> dict:
    """Converte um CSV ou XLSX em fatos canônicos, sem acessar banco de dados."""
    if not isinstance(raw, bytes) or not filename:
        raise ValueError("Arquivo inválido.")
    empresa, origem = _defaults(raw, filename)
    avisos_modelo: list[str] = []
    if filename.lower().endswith(".csv"):
        sheets, headers = _read_csv(raw)
        records = sheets["Lançamentos"]
        is_legacy = LEGADO_MOVIMENTOS <= headers
        if not is_legacy and (not NORMAL_LANCAMENTOS <= headers or not ({"debito", "credito"} <= headers or {"debito_centavos", "credito_centavos"} <= headers)):
            raise ValueError("CSV não corresponde ao modelo normalizado nem ao legado MOVIMENTOS.")
        rows = [row for index, record in enumerate(records, 2) for row in (_legacy_rows(record, empresa, origem) if is_legacy else [_normal_row(record, empresa, origem, index)])]
        accounts = []
    elif filename.lower().endswith(".xlsx"):
        sheets = _read_xlsx(raw)
        xlsx_avisos = sheets.pop("__avisos__", [])
        normal = sheets.get("Lançamentos")
        legacy_records = sheets.get("MOVIMENTOS")
        if normal is not None:
            if not normal or not NORMAL_LANCAMENTOS <= set(normal[0]) or not ({"debito", "credito"} <= set(normal[0]) or {"debito_centavos", "credito_centavos"} <= set(normal[0])):
                raise ValueError("Aba Lançamentos não corresponde ao modelo normalizado.")
            rows = [_normal_row(record, empresa, origem, index) for index, record in enumerate(normal, 2)]
            catalog = sheets.get("Contas", [])
            accounts = _accounts(catalog, False, avisos_modelo)
            records = normal
        elif legacy_records is not None:
            if not legacy_records or not LEGADO_MOVIMENTOS <= set(legacy_records[0]):
                raise ValueError("Aba MOVIMENTOS não corresponde ao modelo legado.")
            dates_by_document = {}
            for record in legacy_records:
                document = _text(record.get("CT2_DOC"), "CT2_DOC", numeric=True)
                if document:
                    dates_by_document.setdefault(document, set()).add(_date(record.get("CT2_DATA")))
            rows, technical = [], []
            for record in legacy_records:
                record = dict(record)
                document = _text(record.get("CT2_DOC"), "CT2_DOC", numeric=True)
                if document and len(dates_by_document[document]) > 1:
                    record["_documento_normalizado"] = f"{document}:{_date(record.get('CT2_DATA'))}"
                if not document:
                    recno = _text(record.get("R_E_C_N_O_"), "R_E_C_N_O_", numeric=True)
                    if recno:
                        technical.append(recno)
                rows.extend(_legacy_rows(record, empresa, origem))
            accounts = _accounts(sheets.get("CONTAS", []), True, avisos_modelo)
            records = legacy_records
        else:
            raise ValueError("XLSX deve conter Lançamentos ou MOVIMENTOS.")
    else:
        raise ValueError("Use um arquivo .xlsx ou .csv.")
    if not rows or len(rows) > MAX_ROWS:
        raise ValueError("Envie de 1 a 5.000 lançamentos.")
    identifiers = [row["linha_id"] for row in rows]
    if len(set(identifiers)) != len(identifiers):
        raise ValueError("Identificador de linha duplicado no arquivo.")
    accounts, avisos_hierarquia = organizar_contas_com_avisos(accounts)
    avisos = locals().get("xlsx_avisos", []) + avisos_modelo + avisos_hierarquia + (["Classificações ausentes permanecem pendentes para revisão."] if any(x["analitica"] and x["grupo"] == "pendente" for x in accounts) else [])
    if filename.lower().endswith(".xlsx") and legacy_records is not None:
        if not any(record.get('APP_ATIVIDADE_CAIXA') for record in records):
            avisos.append("A atividade do ERP não define a atividade da DFC. Complete a classificação de caixa quando necessário.")
        for recno in locals().get("technical", []):
            avisos.append(f"Documento de origem ausente; identificado tecnicamente por RECNO {recno}.")
    if any(isinstance(record.get(field), (int, float)) for record in records for field in ("conta_id", "linha_id", "documento_id")):
        avisos.append("verificarzeros: códigos numéricos foram convertidos sem zeros à esquerda.")
    centros = [{'centro_id': _text(c.get('CTT_CUSTO'), 'CTT_CUSTO'), 'descricao': _text(c.get('CTT_DESC01'), 'CTT_DESC01')} for c in sheets.get('CENCUS', []) if c.get('CTT_CUSTO')]
    from core.dre_importacao import destino_planilha, ler_estrutura, extrair_dre, validar_dre
    for centro, original in zip(centros, (c for c in sheets.get('CENCUS', []) if c.get('CTT_CUSTO'))):
        if 'APP_DRE_DESTINO' in original:
            centro['dre_destino'] = destino_planilha(original['APP_DRE_DESTINO'])
    dre = extrair_dre(ler_estrutura(sheets.get('DRE_GERENCIAL')), accounts, centros)
    if dre and dre['grupos'] is not None:
        validar_dre(dre)
    from core.orcamento_importacao import ler_orcamento, validar_orcamento
    orcamento = ler_orcamento(sheets.get('ORCAMENTO'))
    validar_orcamento(orcamento, accounts)
    origem_linhas = {}
    if filename.lower().endswith('.xlsx') and legacy_records is not None:
        for record in legacy_records:
            doc = _text(record.get('CT2_DOC'), 'CT2_DOC', numeric=True)
            if record.get('APP_DOCUMENTO_ID'): doc = str(record['APP_DOCUMENTO_ID'])
            elif doc and len(dates_by_document.get(doc, [])) > 1: doc += ':' + _date(record.get('CT2_DATA'))
            elif not doc: doc = 'RECNO-' + _text(record.get('R_E_C_N_O_'), 'R_E_C_N_O_', numeric=True)
            origem_linhas.setdefault(doc, []).append(record.get('__linha_excel__'))
    hierarquia = {
        "contas": len(accounts),
        "raizes": sum(not conta["conta_pai_id"] for conta in accounts),
        "sinteticas": sum(not conta["analitica"] for conta in accounts),
        "analiticas": sum(conta["analitica"] for conta in accounts),
        "avisos": avisos_hierarquia,
    }
    return {"rows": rows, "contas": accounts, "avisos": avisos, "centros": centros, "origem_linhas": origem_linhas, "hierarquia": hierarquia, "dre": dre, "orcamento": orcamento}

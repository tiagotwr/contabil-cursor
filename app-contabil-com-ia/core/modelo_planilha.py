"""Modelos XLSX fictícios para importação, inclusive o formato Protheus."""
from __future__ import annotations

import io
from datetime import date

from core.plano_contas import organizar_contas


def _dependencias():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as error:
        raise ValueError("A geração do modelo requer openpyxl.") from error
    return Workbook, Font, PatternFill


def _estilizar(sheet, Font, PatternFill, moeda=()):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF7500")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(x.value or "")) for x in column) + 2, 34)
    for coluna in moeda:
        for cell in sheet[coluna][1:]: cell.number_format = '#,##0.00'


def _anexar_dre(book, contas_ws, centros_ws, dre, Font, PatternFill, normal=False):
    if dre is None:
        return
    destinos = {c['conta_id']:c['destino'] for c in dre['contas']}
    centros = {c['centro_id']:c['destino'] for c in dre['centros']}
    def legivel(valor):
        return 'Prevalece DRE' if valor == '__centro__' else valor
    coluna = contas_ws.max_column + 1
    contas_ws.cell(1,coluna,'dre_destino' if normal else 'APP_DRE_DESTINO')
    for row in range(2,contas_ws.max_row+1):
        contas_ws.cell(row,coluna,legivel(destinos.get(str(contas_ws.cell(row,1).value),'')))
    if centros_ws is None:
        centros_ws = book.create_sheet('CENCUS')
        centros_ws.append(['CTT_CUSTO','CTT_DESC01'])
        for c in dre['centros']:
            centros_ws.append([c['centro_id'],c.get('descricao',c['centro_id'])])
    coluna = centros_ws.max_column + 1
    centros_ws.cell(1,coluna,'APP_DRE_DESTINO')
    for row in range(2,centros_ws.max_row+1):
        centros_ws.cell(row,coluna,legivel(centros.get(str(centros_ws.cell(row,1).value),'')))
    estrutura = book.create_sheet('DRE_GERENCIAL')
    estrutura.append(['codigo','nome','ordem','tipo'])
    for g in dre['grupos']:
        estrutura.append([g[k] for k in ('codigo','nome','ordem','tipo')])
    for sheet in (contas_ws,centros_ws,estrutura):
        _estilizar(sheet,Font,PatternFill)
    for sheet in (contas_ws,centros_ws):
        sheet.column_dimensions[sheet.cell(1,sheet.max_column).column_letter].width = 24
    estrutura.column_dimensions['A'].width = 30
    estrutura.column_dimensions['B'].width = 36


def _anexar_orcamento(book, orcamento, Font, PatternFill):
    if orcamento is None:
        return
    pagina = book.create_sheet("ORCAMENTO")
    pagina.append(["competencia", "conta_id", "valor_dc_centavos"])
    for item in orcamento:
        pagina.append([item["competencia"], item["conta_id"], item["valor_dc_centavos"]])
    _estilizar(pagina, Font, PatternFill)
    pagina.column_dimensions["A"].width = 16
    pagina.column_dimensions["B"].width = 20
    pagina.column_dimensions["C"].width = 24


def gerar_modelo(contas, mov, dre=None, orcamento=None) -> bytes:
    """Modelo normalizado original, mantido para integrações já existentes."""
    Workbook, Font, PatternFill = _dependencias()
    book = Workbook(); lanc = book.active; lanc.title = "Lançamentos"
    headers = ["linha_id", "documento_id", "empresa_id", "origem_id", "data", "conta_id", "debito", "credito", "tipo", "historico", "centro_id", "atividade_caixa", "rubrica_caixa"]
    lanc.append(headers)
    for item in mov:
        values = []
        for key in headers:
            value = item.get(key, "") if key not in {"debito", "credito"} else item.get(key + "_centavos", 0) / 100
            if key == "data" and isinstance(value, str): value = date.fromisoformat(value)
            values.append(value)
        lanc.append(values)
    contas_ws = book.create_sheet("Contas")
    conta_headers = ["conta_id", "descricao", "grupo", "componente", "linha_dre", "conta_pai_id", "ordem", "nivel", "natureza", "classe_bp", "analitica"]
    contas_ws.append(conta_headers)
    for order, account in enumerate(contas, 1):
        contas_ws.append([account["conta_id"], account.get("descricao", ""), account.get("grupo", "pendente"), account.get("componente", ""), account.get("linha_dre", ""), account.get('conta_pai_id',''), account.get('ordem',order), account.get('nivel',1), account.get('natureza',''), account.get('classe_bp','pendente'), account.get("analitica", True)])
    instrucoes = book.create_sheet("Instruções")
    instrucoes.append(["Preencha Lançamentos em reais (até 2 casas). Datas aceitas: data Excel, YYYY-MM-DD, YYYYMMDD ou DD/MM/YYYY."])
    instrucoes.append(["Não altere linha_id: ele evita lançamento duplicado. Códigos devem permanecer texto para preservar zeros à esquerda."])
    _estilizar(lanc, Font, PatternFill, ("G", "H")); _estilizar(contas_ws, Font, PatternFill)
    for cell in lanc["E"][1:]: cell.number_format = "yyyy-mm-dd"
    _anexar_dre(book,contas_ws,None,dre,Font,PatternFill,normal=True)
    _anexar_orcamento(book,orcamento,Font,PatternFill)
    output = io.BytesIO(); book.save(output); return output.getvalue()


def gerar_modelo_protheus(contas, mov, centros=None, dre=None, orcamento=None) -> bytes:
    """Gera MOVIMENTOS/CONTAS/CENCUS no contrato CT2, usando somente dados fornecidos.

    As colunas APP_* preservam classificações explícitas para a evolução do
    importador. CT1_CLASSE usa a convenção da fonte: 1 sintética e 2 analítica;
    CT1_CTASUP repete o pai quando ele existe.
    """
    contas = organizar_contas(contas)
    Workbook, Font, PatternFill = _dependencias()
    book = Workbook(); movimentos = book.active; movimentos.title = "MOVIMENTOS"
    headers = ["CT2_DATA", "CT2_DC", "CT2_DOC", "CT2_LINHA", "CT2_DEBITO", "CT2_CREDIT", "CT2_CCD", "CT2_CCC", "CT2_ITEMD", "CT2_ITEMC", "CT2_ATIVDE", "CT2_ATIVCR", "CT2_HIST", "CT2_VALOR", "CT2_ROTINA", "R_E_C_N_O_", "APP_EMPRESA_ID", "APP_ORIGEM_ID", "APP_DOCUMENTO_ID", "APP_LINHA_DEBITO", "APP_LINHA_CREDITO", "APP_TIPO", "APP_ATIVIDADE_CAIXA", "APP_RUBRICA_CAIXA"]
    movimentos.append(headers)
    documentos = {chave: 900001 + numero for numero, chave in enumerate(sorted({(x.get('documento_id', ''), x['data']) for x in mov}))}
    for recno, item in enumerate(mov, 1):
        debito = int(item.get('debito_centavos', 0)); credito = int(item.get('credito_centavos', 0))
        if bool(debito) == bool(credito):
            raise ValueError("Cada fato do modelo Protheus deve conter somente débito ou crédito positivo.")
        lado_debito = bool(debito); valor_reais = (debito or credito) / 100
        movimentos.append([date.fromisoformat(item['data']), 1 if lado_debito else 2, documentos[(item.get('documento_id', ''), item['data'])], recno,
                            item['conta_id'] if lado_debito else '', item['conta_id'] if not lado_debito else '',
                            item.get('centro_id', '') if lado_debito else '', item.get('centro_id', '') if not lado_debito else '',
                            item.get('rubrica_caixa', '') if lado_debito else '', item.get('rubrica_caixa', '') if not lado_debito else '',
                            item.get('atividade_caixa', '') if lado_debito else '', item.get('atividade_caixa', '') if not lado_debito else '', item.get('historico', ''), valor_reais, 'APP', recno,
                            item.get('empresa_id', ''), item.get('origem_id', ''), item.get('documento_id', ''), item.get('linha_id', '') if lado_debito else '', item.get('linha_id', '') if not lado_debito else '',
                            item.get('tipo', ''), item.get('atividade_caixa', ''), item.get('rubrica_caixa', '')])
    contas_ws = book.create_sheet("CONTAS")
    contas_ws.append(["CT1_CONTA", "CT1_DESC01", "CT1_CLASSE", "CT1_CTASUP", "CT1_RGNV1", "CT1_ZTIPOC", "CT1_ZLINHA", "CT1_ZLDRE1", "R_E_C_N_O_", "APP_GRUPO", "APP_COMPONENTE", "APP_LINHA_DRE", "APP_CONTA_PAI_ID", "APP_ORDEM", "APP_NIVEL", "APP_NATUREZA", "APP_CLASSE_BP", "APP_ANALITICA"])
    for ordem, conta in enumerate(contas, 1):
        contas_ws.append([conta['conta_id'], conta.get('descricao', ''), 2 if conta['analitica'] else 1, conta.get('conta_pai_id', ''), conta.get('natureza', ''), conta.get('componente', ''), conta.get('linha_dre', ''), conta.get('linha_dre', ''), ordem,
                          conta.get('grupo', 'pendente'), conta.get('componente', ''), conta.get('linha_dre', ''), conta.get('conta_pai_id', ''), conta.get('ordem', ordem), conta.get('nivel', ''), conta.get('natureza', ''), conta.get('classe_bp', 'pendente'), conta['analitica']])
    cencus_ws = book.create_sheet("CENCUS")
    cencus_ws.append(["CTT_CUSTO", "CTT_DESC01", "CTT_CRGNV1", "CTT_ZLDRE", "CTT_ZUNEG", "CTT_ZDONO", "R_E_C_N_O_"])
    centros_lista = centros if centros is not None else [{'centro_id': x} for x in sorted({item.get('centro_id', '') for item in mov if item.get('centro_id', '')})]
    for ordem, centro in enumerate(centros_lista, 1):
        codigo = centro.get('centro_id', centro.get('CTT_CUSTO', ''))
        if codigo: cencus_ws.append([codigo, centro.get('descricao', f'Centro {codigo}'), centro.get('hierarquia', ''), centro.get('linha_dre', ''), centro.get('unidade_negocio', ''), centro.get('responsavel', ''), ordem])
    for pagina in (movimentos, contas_ws, cencus_ws): _estilizar(pagina, Font, PatternFill, ("N",) if pagina is movimentos else ())
    for cell in movimentos["A"][1:]: cell.number_format = "yyyy-mm-dd"
    for codigo_coluna in ("E", "F", "G", "H", "I", "J", "Q", "R"):
        for cell in movimentos[codigo_coluna][1:]: cell.number_format = "@"
    for pagina in (contas_ws, cencus_ws):
        for cell in pagina["A"][1:]: cell.number_format = "@"
    _anexar_dre(book,contas_ws,cencus_ws,dre,Font,PatternFill)
    _anexar_orcamento(book,orcamento,Font,PatternFill)
    output = io.BytesIO(); book.save(output); return output.getvalue()
